# Verifica e instala o openpyxl se não estiver instalado
try:
    import openpyxl
    import re
except ImportError:
    import subprocess
    import sys
    print("Pacote 'openpyxl' não encontrado. Instalando...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    

from openpyxl.utils import get_column_letter

# Caminho para seu arquivo
CAMINHO_ARQUIVO = r'Z:\Cloud\OneDrive\Desktop\0 ELLO\py\data.xlsx'  # <-- Altere aqui

# Abre a planilha
wb = openpyxl.load_workbook(CAMINHO_ARQUIVO)
geral = wb['GERALPY']

# Cabeçalho da tabela
headers = ["Worksheet", "Nome", "Função", "Salario Base", "Adicional Noturno 40%", "Adicional Insalubridade", "DSR Adicional Noturno"]
for col, header in enumerate(headers, 1):
    geral.cell(row=1, column=col).value = header

# Para cada worksheet Table 1 até Table 147
for i in range(1, 148):
    sheet_name = f'Table {i}'
    if sheet_name not in wb.sheetnames:
        continue

    ws = wb[sheet_name]
    nome_ref = funcao_ref = salario_ref = ''
    adicionais = {
        "Adicional Noturno 40%": "0",
        "Adicional Insalubridade": "0",
        "DSR Adicional Noturno": "0"
    }

    # Busca Nome e Função
    for row in range(1, 100):
        for col in range(1, 20):
            val = ws.cell(row=row, column=col).value
            if isinstance(val, str) and 'nome' in val.lower():
                nome_ref = f"='{sheet_name}'!{get_column_letter(col)}{row+1}"
                funcao_ref = f"='{sheet_name}'!{get_column_letter(col)}{row+2}"

    # Busca Salario Base mesmo com célula mesclada e valores múltiplos
    salario_ref = "0"
    for merged_range in ws.merged_cells.ranges:
        top_left = merged_range.min_row, merged_range.min_col
        val = ws.cell(*top_left).value
        if isinstance(val, str) and 'salario base' in val.strip().lower():
            row_below = merged_range.max_row + 1
            col_below = merged_range.min_col
            cell_value = ws.cell(row=row_below, column=col_below).value
            col_letter = get_column_letter(col_below)
            cell_ref = f"'{sheet_name}'!{col_letter}{row_below}"

            if isinstance(cell_value, str):
                numeros = re.findall(r'\d{1,3}(?:\.\d{3})*,\d{2}', cell_value)
                if len(numeros) > 1:
                    # Usar função MID (em inglês) e com INDIRETO para evitar quebra
                    salario_ref = f'=MID(INDIRECT("{cell_ref}"),1,8)'
                else:
                    salario_ref = f"='{sheet_name}'!{col_letter}{row_below}"
            else:
                salario_ref = f"='{sheet_name}'!{col_letter}{row_below}"
            break


    # Busca Descrição e Vencimentos
    desc_row = None
    for row in range(1, 100):
        for col in range(1, 10):
            if ws.cell(row=row, column=col).value == "Descrição":
                desc_row = row
                desc_col = col
                break
        if desc_row:
            break

    if desc_row:
        # Localiza coluna Vencimentos
        venc_col = None
        for col in range(1, 20):
            if ws.cell(row=desc_row, column=col).value == "Vencimentos":
                venc_col = col
                break

        # Percorre linhas com os adicionais
        if venc_col:
            for row in range(desc_row + 1, desc_row + 30):
                key = ws.cell(row=row, column=desc_col).value
                if isinstance(key, str):
                    for adicional in adicionais.keys():
                        if key.strip().lower() == adicional.lower():
                            adicionais[adicional] = f"='{sheet_name}'!{get_column_letter(venc_col)}{row}"

    # Grava os dados na GERALPY
    geral.cell(row=i + 1, column=1).value = i  # número da worksheet
    geral.cell(row=i + 1, column=2).value = nome_ref or "0"
    geral.cell(row=i + 1, column=3).value = funcao_ref or "0"
    geral.cell(row=i + 1, column=4).value = salario_ref or "0"
    geral.cell(row=i + 1, column=5).value = adicionais["Adicional Noturno 40%"]
    geral.cell(row=i + 1, column=6).value = adicionais["Adicional Insalubridade"]
    geral.cell(row=i + 1, column=7).value = adicionais["DSR Adicional Noturno"]

# Salva o arquivo
wb.save(CAMINHO_ARQUIVO)
print("Planilha GERALPY atualizada com referências!")
